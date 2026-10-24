
import re
import tempfile
import logging
from pathlib import Path

import fitz
from docx import Document
from flask import Flask, render_template, request
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

logging.basicConfig(level=logging.INFO)

ALLOWED_EXTENSIONS = {"pdf", "docx"}

SKILLS = [
    "python", "java", "javascript", "html", "css",
    "react", "flask", "django", "php", "sql", "mysql",
    "aws", "ec2", "s3", "iam", "linux", "windows",
    "networking", "docker", "kubernetes", "git", "github",
    "devops", "ci/cd", "machine learning", "data analysis",
    "pandas", "numpy", "excel", "communication",
    "problem solving", "rest api", "jenkins"
]


def allowed_file(filename):
    if not filename or "." not in filename:
        return False

    extension = filename.rsplit(".", 1)[1].lower()
    return extension in ALLOWED_EXTENSIONS


def extract_text(file_path, extension):
    """Extract readable text from PDF or DOCX."""
    if extension == "pdf":
        with fitz.open(file_path) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    document = Document(file_path)
    parts = [paragraph.text for paragraph in document.paragraphs]

    # Also read text from DOCX tables.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)

    return "\n".join(parts)


def contains_keyword(text, keyword):
    """Match a complete keyword, ignoring letter case."""
    pattern = (
        r"(?<!\w)" + re.escape(keyword) + r"(?!\w)"
    )
    return re.search(pattern, text, flags=re.IGNORECASE) is not None


def find_skills(text):
    return [
        skill for skill in SKILLS
        if contains_keyword(text, skill)
    ]


def analyze_resume(resume_text, job_description):
    resume_skills = find_skills(resume_text)
    required_skills = find_skills(job_description)

    matched = [
        skill for skill in required_skills
        if skill in resume_skills
    ]
    missing = [
        skill for skill in required_skills
        if skill not in resume_skills
    ]

    # The score measures detected skill keyword overlap only.
    score = (
        round(len(matched) / len(required_skills) * 100)
        if required_skills else 0
    )

    email_match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        resume_text
    )

    phone_match = re.search(
        r"(?<!\w)(?:\+91[\s-]?)?[6-9]\d{9}(?!\w)",
        resume_text
    )

    suggestions = []

    if not email_match:
        suggestions.append(
            "Check whether your resume includes a contact email."
        )

    if missing:
        suggestions.append(
            "If you genuinely have these skills, highlight relevant "
            "experience or projects: " + ", ".join(missing)
        )

    if not resume_skills:
        suggestions.append(
            "Consider adding a clearly labelled skills section."
        )

    if not suggestions:
        suggestions.append(
            "Review your project descriptions and ensure your skills "
            "are supported by real examples."
        )

    return {
        "score": score,
        "resume_skills": resume_skills,
        "matched_skills": matched,
        "missing_skills": missing,
        "suggestions": suggestions,
        "email": email_match.group(0) if email_match else "Not detected",
        "phone": phone_match.group(0) if phone_match else "Not detected",
        "word_count": len(resume_text.split()),
        "required_count": len(required_skills)
    }


@app.route("/", methods=["GET", "POST"])
def index():
    results = None
    error = None

    if request.method == "POST":
        uploaded_file = request.files.get("resume")
        job_description = request.form.get(
            "job_description", ""
        ).strip()

        if not uploaded_file or not uploaded_file.filename:
            error = "Please choose a resume file."

        elif not allowed_file(uploaded_file.filename):
            error = "Only PDF and DOCX files are supported."

        elif not job_description:
            error = "Please enter the job description."

        elif not find_skills(job_description):
            error = (
                "No supported skills were found in the job description. "
                "Try adding skills such as Python, AWS, Linux or SQL."
            )

        else:
            safe_name = secure_filename(uploaded_file.filename)

            if not safe_name:
                error = "The uploaded filename is invalid."
            else:
                extension = safe_name.rsplit(".", 1)[1].lower()
                temp_path = None

                try:
                    # A temporary file is used instead of retaining resumes.
                    with tempfile.NamedTemporaryFile(
                        suffix="." + extension,
                        dir=UPLOAD_DIR,
                        delete=False
                    ) as temp_file:
                        temp_path = Path(temp_file.name)
                        uploaded_file.save(temp_file)

                    resume_text = extract_text(temp_path, extension)

                    if not resume_text.strip():
                        error = (
                            "No readable text was found. "
                            "Scanned PDFs may require OCR."
                        )
                    else:
                        results = analyze_resume(
                            resume_text, job_description
                        )

                except Exception:
                    app.logger.exception("Resume processing failed")
                    error = (
                        "Unable to read this document. Please check "
                        "that it is a valid PDF or DOCX file."
                    )

                finally:
                    if temp_path and temp_path.exists():
                        temp_path.unlink()

    return render_template(
        "index.html",
        results=results,
        error=error
    )


@app.errorhandler(413)
def file_too_large(_error):
    return render_template(
        "index.html",
        results=None,
        error="The file is too large. Maximum size is 5 MB."
    ), 413


if __name__ == "__main__":
    app.run(debug=True)
